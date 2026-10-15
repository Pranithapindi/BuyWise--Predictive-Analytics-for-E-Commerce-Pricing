"""
Export API Router — export price history as CSV, Excel, or PDF.
"""
import io
import csv
from datetime import datetime
from fastapi import APIRouter, Depends, HTTPException, Query, status
from fastapi.responses import StreamingResponse
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.dependencies import get_db, get_current_active_user
from app.database.models import User
from app.database.repositories.product_repo import ProductRepository
from app.database.repositories.price_history_repo import PriceHistoryRepository

router = APIRouter()


@router.get("/{product_id}")
async def export_data(
    product_id: str,
    format: str = Query("csv", pattern="^(csv|excel|pdf)$"),
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_active_user),
):
    """Export product price history as CSV, Excel, or PDF."""
    product_repo = ProductRepository(db)
    product = await product_repo.get_by_id_and_user(product_id, current_user.id)
    if not product:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Product not found")

    history_repo = PriceHistoryRepository(db)
    records = await history_repo.get_by_product(product_id, limit=5000)

    product_name = (product.name or "product")[:30].replace(" ", "_")
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    filename_base = f"price_history_{product_name}_{timestamp}"

    # ── CSV ───────────────────────────────────────────────────────────────────
    if format == "csv":
        output = io.StringIO()
        writer = csv.writer(output)
        writer.writerow([
            "Date", "Price (₹)", "Original Price (₹)", "Discount (%)",
            "Rating", "Reviews", "Availability", "Seller", "Delivery (₹)"
        ])
        for r in records:
            writer.writerow([
                r.scraped_at.strftime("%Y-%m-%d %H:%M"),
                float(r.price),
                float(r.original_price) if r.original_price else "",
                float(r.discount_pct) if r.discount_pct else "",
                float(r.rating) if r.rating else "",
                r.review_count or "",
                r.availability or "",
                r.seller or "",
                float(r.delivery_charges) if r.delivery_charges else "",
            ])
        output.seek(0)
        return StreamingResponse(
            iter([output.getvalue()]),
            media_type="text/csv",
            headers={"Content-Disposition": f'attachment; filename="{filename_base}.csv"'},
        )

    # ── Excel ─────────────────────────────────────────────────────────────────
    elif format == "excel":
        try:
            import openpyxl
            from openpyxl.styles import Font, PatternFill, Alignment
        except ImportError:
            raise HTTPException(status_code=500, detail="openpyxl not installed")

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Price History"

        headers = [
            "Date", "Price (₹)", "Original Price (₹)", "Discount (%)",
            "Rating", "Reviews", "Availability", "Seller", "Delivery (₹)"
        ]
        header_fill = PatternFill("solid", fgColor="6366F1")
        header_font = Font(color="FFFFFF", bold=True)

        for col_idx, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")
            ws.column_dimensions[cell.column_letter].width = 18

        for row_idx, r in enumerate(records, 2):
            ws.append([
                r.scraped_at.strftime("%Y-%m-%d %H:%M"),
                float(r.price),
                float(r.original_price) if r.original_price else None,
                float(r.discount_pct) if r.discount_pct else None,
                float(r.rating) if r.rating else None,
                r.review_count,
                r.availability,
                r.seller,
                float(r.delivery_charges) if r.delivery_charges else None,
            ])

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return StreamingResponse(
            iter([output.read()]),
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f'attachment; filename="{filename_base}.xlsx"'},
        )

    # ── PDF ───────────────────────────────────────────────────────────────────
    elif format == "pdf":
        try:
            from reportlab.lib.pagesizes import A4, landscape
            from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
            from reportlab.lib.styles import getSampleStyleSheet
            from reportlab.lib import colors
        except ImportError:
            raise HTTPException(status_code=500, detail="reportlab not installed")

        output = io.BytesIO()
        doc = SimpleDocTemplate(output, pagesize=landscape(A4))
        styles = getSampleStyleSheet()
        elements = []

        # Title
        elements.append(Paragraph(f"Price History: {product.name or 'Product'}", styles["Title"]))
        elements.append(Paragraph(f"Website: {product.website} | Exported: {datetime.now().strftime('%Y-%m-%d')}", styles["Normal"]))
        elements.append(Spacer(1, 12))

        # Table data
        table_data = [["Date", "Price (₹)", "Original (₹)", "Disc%", "Rating", "Reviews", "Availability"]]
        for r in records[:500]:  # Limit PDF to 500 rows
            table_data.append([
                r.scraped_at.strftime("%Y-%m-%d"),
                f"₹{float(r.price):.0f}",
                f"₹{float(r.original_price):.0f}" if r.original_price else "-",
                f"{float(r.discount_pct):.0f}%" if r.discount_pct else "-",
                str(float(r.rating)) if r.rating else "-",
                str(r.review_count) if r.review_count else "-",
                r.availability or "-",
            ])

        t = Table(table_data, repeatRows=1)
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#6366F1")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")]),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#E2E8F0")),
            ("FONTSIZE", (0, 0), (-1, -1), 9),
        ]))
        elements.append(t)
        doc.build(elements)
        output.seek(0)
        return StreamingResponse(
            iter([output.read()]),
            media_type="application/pdf",
            headers={"Content-Disposition": f'attachment; filename="{filename_base}.pdf"'},
        )
